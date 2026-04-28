"""User-uploaded research notes / spreadsheets. Extract plaintext at upload
time and stash both the original file (under reports/<id>/uploads/) and the
extracted text (in the DB) so analyst agents can pull from it as a tool."""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

from sqlmodel import Session, select

from api.db import engine
from api.models import UploadedDocument
from api.settings import settings

log = logging.getLogger("uploads")

# Per-document storage cap. Anything larger is truncated -- agents don't need
# 5MB of extracted text in their context window, and Postgres TEXT will store
# this comfortably.
MAX_EXTRACTED_CHARS = 200_000

# Per-row preview shown in the listing tool so the agent can decide what to
# dive into.
SUMMARY_CHARS = 600


def uploads_dir(report_id: int) -> Path:
    """Where uploaded originals live. Sibling of reports/<id>/charts/."""
    d = settings.reports_dir / str(report_id) / "uploads"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader  # imported lazily so tests without pypdf still pass

    reader = PdfReader(io.BytesIO(data))
    chunks: list[str] = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - some encrypted/corrupt PDFs throw
            continue
    return "\n\n".join(c.strip() for c in chunks if c.strip())


def _extract_csv(data: bytes) -> str:
    """Render a CSV/TSV as a compact markdown table-ish preview. We don't try
    to be clever about delimiters -- csv.Sniffer handles the common cases."""
    text = data.decode("utf-8", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        return ""
    out_lines = [" | ".join(rows[0])]
    out_lines.append(" | ".join("---" for _ in rows[0]))
    for r in rows[1:]:
        out_lines.append(" | ".join(r))
    return "\n".join(out_lines)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook  # imported lazily

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            parts.append(" | ".join("" if v is None else str(v) for v in row))
        parts.append("")
    return "\n".join(parts)


def extract_text(filename: str, mime: str, data: bytes) -> str:
    """Dispatch on filename suffix + mime type. Always returns plaintext;
    truncates to MAX_EXTRACTED_CHARS so a single 1000-page upload can't blow
    out an analyst's context window."""
    name = filename.lower()
    try:
        if name.endswith(".pdf") or "pdf" in mime:
            text = _extract_pdf(data)
        elif name.endswith((".csv", ".tsv")) or mime in {"text/csv", "text/tab-separated-values"}:
            text = _extract_csv(data)
        elif name.endswith((".xlsx", ".xlsm")):
            text = _extract_xlsx(data)
        elif name.endswith((".md", ".txt", ".markdown")) or mime.startswith("text/"):
            text = data.decode("utf-8", errors="replace")
        else:
            # Best-effort: try utf-8. Worst case the agent sees gibberish and
            # ignores the doc, which beats hard-failing the upload.
            text = data.decode("utf-8", errors="replace")
    except Exception as e:  # noqa: BLE001 - extraction never fails the upload
        log.warning("extraction failed for %s (%s): %s", filename, mime, e)
        return f"[extraction failed: {e}]"

    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS] + "\n\n[... truncated ...]"
    return text.strip()


def save_upload(report_id: int, filename: str, mime: str, data: bytes) -> UploadedDocument:
    """Write the original to disk, extract text, and persist a row. Idempotent
    on (report_id, filename) -- re-uploading the same name overwrites both
    the file and the DB row."""
    target = uploads_dir(report_id) / filename
    target.write_bytes(data)
    text = extract_text(filename, mime, data)
    summary = text[:SUMMARY_CHARS] + ("…" if len(text) > SUMMARY_CHARS else "")

    with Session(engine) as session:
        existing = session.exec(
            select(UploadedDocument)
            .where(UploadedDocument.report_id == report_id)
            .where(UploadedDocument.filename == filename)
        ).first()
        if existing:
            existing.mime = mime
            existing.size_bytes = len(data)
            existing.extracted_text = text
            existing.summary = summary
            session.add(existing)
            session.commit()
            session.refresh(existing)
            return existing
        doc = UploadedDocument(
            report_id=report_id,
            filename=filename,
            mime=mime,
            size_bytes=len(data),
            extracted_text=text,
            summary=summary,
        )
        session.add(doc)
        session.commit()
        session.refresh(doc)
        return doc


def list_documents(report_id: int) -> list[UploadedDocument]:
    with Session(engine) as session:
        return list(
            session.exec(
                select(UploadedDocument)
                .where(UploadedDocument.report_id == report_id)
                .order_by(UploadedDocument.created_at)  # type: ignore[arg-type]
            ).all()
        )


def get_document(report_id: int, filename: str) -> UploadedDocument | None:
    with Session(engine) as session:
        return session.exec(
            select(UploadedDocument)
            .where(UploadedDocument.report_id == report_id)
            .where(UploadedDocument.filename == filename)
        ).first()


def delete_document(report_id: int, filename: str) -> bool:
    doc = get_document(report_id, filename)
    if not doc:
        return False
    with Session(engine) as session:
        row = session.get(UploadedDocument, doc.id)
        if row:
            session.delete(row)
            session.commit()
    path = uploads_dir(report_id) / filename
    if path.exists():
        path.unlink()
    return True
